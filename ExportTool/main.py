# coding=UTF-8
import sys

import docx
import re
#import xlrd
import xlwt
#from xlutils.copy import copy


class QuestionData:
    id = 0
    qus = "题目"
    opt = []
    ans = "答案"
    analysis = "解析"

    def __init__(self, ques):
        quesarr = ques.split(".")
        self.id = int(quesarr[0])
        self.qus = quesarr[1]
        self.opt = []
        self.ans = "答案"
        self.analysis = "解析"

    def AddOption(self, opts):
        optstr = opts.replace(" ", "").replace("\t","")
        # optstr = opts.replace("\t", "")
        optstr = optstr.replace("A.", "A.")
        optstr = optstr.replace("B.", "_B.")
        optstr = optstr.replace("C.", "_C.")
        optstr = optstr.replace("D.", "_D.")
        arr = optstr.split("_")
        # self.opt = self.opt + arr
        for a in arr:
            if a != '' and len(a) > 0 :
                self.opt.append(a)
        if len(self.opt) >4 :
            print( "id:%d 选项格式有错(或者下一题的题目格式有错,题目一定是数字+.的格式)，请检查,选项超过4个 %s" % (self.id, " ".join(self.opt) ))
        self.opt.sort()

    def AddAnswer(self, answer):
        self.ans = answer

    def AddAnalysis(self,analysis):
        self.analysis =  analysis

    def tostring(self):
        return str(self.id) + "." + self.qus + " 选项:" + " ".join(self.opt) + " 答案:" + self.ans ;


def main():
    if len(sys.argv) > 1:
        fileName = sys.argv[1]
    else:
        fileName = "files/宋元400题目.docx"
    doc = getDoc(fileName)
    global questionRule
    global answerRole
    global optionROle
    global analysisRole
    questionRule = re.compile("^[0-9]*\.[^答案]")
    answerRole = re.compile("^[0-9]*\.答案")
    optionROle = re.compile("^[A-Z].")
    analysisRole = re.compile("^解析：")
    questions = []
    questionsDic = {}
    isInAnalysis = False
    for d in doc.paragraphs:
        str = d.text.replace("．", ".").replace("\t","")
        if isQustion(str):  # 问题
            curQu = QuestionData(str)
            questions.append(curQu)
            questionsDic[curQu.id]=curQu
        elif isOption(str):  # 选项
            curQu.AddOption(str)
        elif isAnswer(str):  # 答案
            anarr = str.split(".")
            anId = int(anarr[0])
            tmpanswer = str[-1]
            analysisArr = []
            if questionsDic.__contains__(anId):
                tmpQu = questionsDic[anId]
                if tmpQu:
                    tmpQu.AddAnswer(tmpanswer)
        elif ( isAnalysis(str) or isInAnalysis ) :#解析
            str = str.replace("解析:","").replace("解析：","")
            isInAnalysis = True
            if anId > 0:
                analysisArr.append(str)
            if str[-1] == "。" :
                isInAnalysis = False
                if tmpQu:
                    tmpQu.AddAnalysis("".join(analysisArr))

    fileName = fileName.split(".")[0] + ".xlsx"
    writeToExcel(fileName, "题目",questions)



def isAnalysis(str):
    global analysisRole
    if analysisRole.match(str):
        return True
    return False

def isAnswer(str):
    global answerRole
    if answerRole.match(str):
        return True
    return False


def isOption(str):
    global optionROle
    if optionROle.match(str):
        return True
    return False


def isQustion(str):
    global questionRule
    if questionRule.match(str):
        return True
    return False


def getDoc(filename):
    d = docx.Document(filename)
    return d

def writeToExcel(path,sheet_name,questionDatas):
    index = len(questionDatas) ##读取所需要的行
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet(sheet_name)
    j = 0
    for i in range(0,index):
        que = questionDatas[i]
        sheet.write(i, j, que.id)
        j = j+1
        sheet.write(i,j,que.qus)
        j = j+1
        optLen = len(que.opt)
        if optLen < 6 :
            optLen = 6
        for n in range(j,j+len(que.opt)):
            tmpOp=que.opt[n-j]
            tmpOp = tmpOp.replace("A.","").replace("B.","").replace("C.","").replace("D.","")
            sheet.write(i, n, tmpOp)
        j = j + optLen
        sheet.write(i,j,que.ans)
        sheet.write(i,j+1,que.analysis)
        j = 0
    workbook.save(path)



if __name__ == '__main__':
    main()
