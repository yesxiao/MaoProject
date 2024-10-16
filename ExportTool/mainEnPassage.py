# coding=UTF-8
import sys
import docx
import re
# import xlrd
# import xlwt
# from xlutils.copy import copy
import xlwt as xlwt


class QuestionData:
    id = 0
    oldId = 0
    qus = "题目"
    ansIds = None
    ans = None
#添加答案
    def AddAns(self,ans):
        if self.ans != None:
            print("id=%d 重复赋值" % self.oldId)
            return
        if len(ans) > 0 :
            if ans[0] == '':
                del ans[0]
        for i in range(0,len(ans)):
            ans[i] = ans[i].strip()
        self.ans = ans

    def AddQue(self,qus):
        qus = re.sub("._+",".______",qus)
        self.qus = qus
        res = re.findall("[0-9]\s*[.]_",qus)
        #res.sort()
        for i in range(0,len(res)):
            s = res[i]
            s = str.strip(s)
            s = s.replace(" ","")
            res[i]=s
        self.ansIds = res

    def isValid(self):
        if self.ans == None:
            print("id=%d 答案不存在" % self.oldId )
            return False
        if self.ansIds == None:
            print("id=%d 答案不存在或格式不正确" % self.oldId )
            return False
        isSameParamsNum = len(self.ans) == len(self.ansIds)
        if not isSameParamsNum:
            print("id=%d 题目和答案的数量不正确,题目 %d 个，答案 %d 个" % (self.oldId , len(self.ansIds) , len(self.ans)) )
        return isSameParamsNum

def main():
    if len(sys.argv) > 1:
        fileName = sys.argv[1]
    else:
        fileName = "files/14-高中英语-14综合语法填空"
    if not fileName.endswith(".docx"): # 文件夹
        for path in os.listdir(fileName):
            if path.endswith(".docx"):
                path = fileName + "/" + path
                HandleFile(path)
    else:
        HandleFile(fileName)

def HandleFile(fileName):
    # if len(sys.argv) > 1:
    #     fileName = sys.argv[1]
    # else:
    #     fileName = "files/14-高中英语-14综合语法填空.docx"
    doc = getDoc(fileName)
    questions = []
    questionsDic = {}
    isStartQue = False
    isStartAns = False
    hasError = False
    curQu = None
    isInQue = False
    isInAns = False
    queArr = []
    ansArr = []
    tmpId = 0
    for d in doc.paragraphs:
        str = d.text
        if str.find("//题目开始") != -1:
            isStartQue = True
            continue
        elif str.find("//题目结束") != -1:
            isStartQue = False
            isInQue = False
            #为最后一个设置题目
            if curQu and not questionsDic.__contains__(curQu.id):
                curQu.AddQue("\n".join(queArr))
                questionsDic[curQu.id] = curQu
                questions.append(curQu)
            continue
        elif str.find("//答案开始") != -1:
            isStartAns = True
            continue
        elif str.find("答案结束") != -1:
            addAns(ansArr, tmpId, questionsDic)
            isStartAns = False
            break
        if isStartQue and isStartAns:
            print("严重错误，题目还未结束，答案已经开始，请检查相关标记")
            hasError = True
            break
        if isStartQue:
            if str.lstrip().startswith("Passage"):
                if curQu and not questionsDic.__contains__(curQu.oldId):
                    curQu.AddQue("\n".join(queArr))
                    questionsDic[curQu.oldId] = curQu
                    questions.append(curQu)
                isInQue = True
                queArr = []
                curQu = QuestionData()
                curQu.id = len(questions) + 1
                curQu.oldId = int(str.strip().replace("Passage",""))
                continue
            elif isInQue:
                queArr.append(str)
        if isStartAns:
            if str.lstrip().startswith("Passage"):
                addAns(ansArr,tmpId,questionsDic)
                # if ansArr and len(ansArr) > 0 :
                #     an = "\n".join(ansArr)
                #     arr = re.split("[0-9]\s*[.]",an)
                #     if not questionsDic.__contains__(tmpId):
                #         print("id=%d 只有答案，没有题目" % tmpId )
                #     else:
                #         curQu = questionsDic[tmpId]
                #         curQu.AddAns(arr)
                ansArr = []
                isInAns = True
                try:
                    tmpId = int(str.strip().replace("Passage", ""))
                except :
                    print( "解析错误:" + str )
                continue
            if isInAns:
                ansArr.append(str.strip())

    if not hasError:
        fileName = fileName.split(".")[0] + ".xlsx"
        writeToExcel(fileName, "题目", questions)

def addAns(ansArr,tmpId,questionsDic):
    if ansArr and len(ansArr) > 0:
        an = "\n".join(ansArr)
        arr = re.split("[0-9]+\s*[.]", an)
        arr.pop(0)
        for i in range(0, len(arr)):
            s = arr[i]
            s = str.strip(s)
            arr[i] = s
        if not questionsDic.__contains__(tmpId):
            print("id=%d 只有答案，没有题目" % tmpId)
        else:
            curQu = questionsDic[tmpId]
            curQu.AddAns(arr)


def getDoc(filename):
    d = docx.Document(filename)
    return d


def writeToExcel(path, sheet_name, questionDatas):
    index = len(questionDatas)  ##读取所需要的行
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet(sheet_name)
    j = 0
    for i in range(0, index):
        que = questionDatas[i]
        if not que.isValid():
            continue
        sheet.write(i, j, que.id)
        sheet.col(i).width = 200*20
        sheet.col(i).height = 300*20
        j = j + 1
        sheet.write(i, j, que.qus)
        j = j + 1
        ansLen = len(que.ans)
        for n in range(j, j + ansLen):
            tmpOp = que.ans[n - j]
            tmpOp = tmpOp.replace("A.", "").replace("B.", "").replace("C.", "").replace("D.", "").replace("E.","").replace("F.","")
            sheet.write(i, n, tmpOp)
        j = 0
    workbook.save(path)


if __name__ == '__main__':
    main()
